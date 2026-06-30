#!/usr/bin/env python
"""A focused 3-slide deck:
  1) what a multi-copy gene family IS in our definition (read-conflict graph),
  2) a table of methods to detect multi-copy gene families (the gathered related-work),
  3) how a runaway read overlapping the last exon is prevented from wrongly extending the intron chain.

Output: bench/slides/three_slides.pptx (+ PNGs). Run: /home/juanfra/miniforge3/bin/python bench/make_3slide_deck.py
"""
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import networkx as nx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = os.path.join(os.path.dirname(__file__), "slides")
os.makedirs(OUT, exist_ok=True)
plt.rcParams.update({"font.size": 15, "font.family": "DejaVu Sans", "axes.spines.top": False,
                     "axes.spines.right": False, "figure.dpi": 150})
BLUE, RED, GREEN, GREY, NAVY = "#377eb8", "#e41a1c", "#4daf4a", "#999999", (20, 40, 90)


def savefig(name):
    p = os.path.join(OUT, name)
    plt.savefig(p, bbox_inches="tight", facecolor="white"); plt.close(); return p


def fig_conflict():
    fig, ax = plt.subplots(figsize=(9.5, 4.4))
    G = nx.Graph()
    fam1, fam2, single, dom = ["A0", "A1", "A2"], ["B0", "B1"], ["S"], ["D0", "D1"]
    G.add_nodes_from(fam1 + fam2 + single + dom)
    G.add_edges_from([("A0", "A1"), ("A1", "A2"), ("A0", "A2"), ("B0", "B1")])
    pos = {"A0": (0, 1), "A1": (1, 1.6), "A2": (1, .4), "B0": (3, 1.3), "B1": (4, .8),
           "S": (5.7, 1.1), "D0": (3.3, -.7), "D1": (4.3, -.7)}
    nx.draw_networkx_nodes(G, pos, nodelist=fam1, node_color=BLUE, node_size=1500, ax=ax)
    nx.draw_networkx_nodes(G, pos, nodelist=fam2, node_color=GREEN, node_size=1500, ax=ax)
    nx.draw_networkx_nodes(G, pos, nodelist=single + dom, node_color=GREY, node_size=1500, ax=ax)
    nx.draw_networkx_edges(G, pos, width=3, ax=ax)
    nx.draw_networkx_labels(G, pos, font_color="white", font_weight="bold", font_size=12, ax=ax)
    ax.text(0.5, 2.15, "family 1\n(tandem array)", ha="center", color=BLUE, fontsize=12, weight="bold")
    ax.text(3.5, 1.85, "family 2\n(cross-chrom pair)", ha="center", color=GREEN, fontsize=12, weight="bold")
    ax.text(5.7, 1.7, "single-copy\n(no tie → not a family)", ha="center", fontsize=11)
    ax.text(3.8, -1.25, "domain-sharers\n(share 1 exon, NO tie → not merged)", ha="center", fontsize=11)
    ax.text(2.0, -2.05, "edge = a read fits BOTH loci comparably (a 'tie')   •   family = connected component",
            ha="center", fontsize=12.5, style="italic")
    ax.axis("off"); ax.set_ylim(-2.35, 2.6); ax.set_xlim(-0.8, 6.6)
    return savefig("s_conflict.png")


def fig_runaway():
    fig, ax = plt.subplots(figsize=(11, 5.2)); ax.axis("off")
    # reference transcript: 3 exons (boxes) joined by introns (lines). terminal = exon 3.
    exons = [(1.0, 2.2), (3.4, 4.4), (5.6, 6.6)]  # (x0,x1) of each exon
    ey = 4.0
    for i, (x0, x1) in enumerate(exons):
        ax.add_patch(plt.Rectangle((x0, ey), x1 - x0, 0.45, fc=BLUE, ec="black"))
        ax.text((x0 + x1) / 2, ey + 0.22, f"e{i+1}", ha="center", va="center", color="white", weight="bold")
    for (a, b) in [(exons[0][1], exons[1][0]), (exons[1][1], exons[2][0])]:
        ax.plot([a, b], [ey + 0.22, ey + 0.22], color=BLUE, lw=1.5)
    ax.text(3.8, ey + 0.75, "real transcript:  intron chain = (e1|e2, e2|e3),  terminal exon = e3", ha="center", fontsize=12, weight="bold")
    # 4 good reads agreeing
    for k, y in enumerate([3.3, 3.0, 2.7, 2.4]):
        ax.add_patch(plt.Rectangle((1.0, y), 5.6, 0.13, fc=GREY, ec="black"))
    ax.text(0.9, 2.85, "4 reads agree", ha="right", fontsize=11, color="black")
    # runaway read: overshoots e3 AND splices into a phantom downstream exon
    yr = 1.7
    ax.add_patch(plt.Rectangle((1.0, yr), 5.6, 0.16, fc=RED, ec="black"))
    ax.plot([6.6, 7.8], [yr + 0.08, yr + 0.08], color=RED, lw=1.5, ls="--")  # phantom intron
    ax.add_patch(plt.Rectangle((7.8, yr), 1.4, 0.16, fc=RED, ec="black"))    # phantom downstream exon
    ax.text(0.9, yr + 0.08, "runaway read", ha="right", fontsize=11, color=RED, weight="bold")
    ax.annotate("read-through / chimera / mis-clip:\novershoots e3 + adds a PHANTOM junction", (8.5, yr + 0.4),
                (6.7, 0.9), fontsize=11, color=RED, ha="left", arrowprops=dict(arrowstyle="->", color=RED))
    # the two defenses
    ax.add_patch(plt.Rectangle((0.4, -1.45), 9.6, 1.85, fc="#f4f4f4", ec=GREY))
    ax.text(0.6, 0.1, "Two defenses (de-novo assembly, denovo_assemble.rs):", fontsize=12.5, weight="bold", color="#14285a")
    ax.text(0.6, -0.35, "①  EXACT intron-chain key — the runaway's chain (extra phantom junction) ≠ the real chain →\n"
            "     it forms its OWN skeleton; a lone runaway fails the ≥3-read quorum → DROPPED. It never extends the real chain.",
            fontsize=11.5, va="top")
    ax.text(0.6, -1.0, "②  Terminal boundary = the k-th-SUPPORTED end (k=2), not the union (max-end) → a single read's\n"
            "     overshoot of e3 is trimmed to the 2nd-furthest position. Verified: pass1_robust_trims_a_runaway_terminal_read.",
            fontsize=11.5, va="top")
    ax.set_xlim(-1.2, 10.4); ax.set_ylim(-1.6, 5.0)
    return savefig("s_runaway.png")


METHODS = [
    ("Method", "Signal / level", "Approach", "Limit on RECENT near-identical paralogs"),
    ("OrthoFinder", "protein homology", "DIAMOND all-vs-all + MCL + gene trees", "near-identical copies collapse / over-merge"),
    ("Pfam / InterPro", "domain HMM", "group by shared domain architecture", "groups domain-sharers (FP); domain ≠ copy"),
    ("Ensembl Compara", "phylogenetic", "gene-tree ↔ species-tree reconciliation", "fails on recent near-identical paralogs"),
    ("Mash / minimizer-Jaccard", "alignment-free", "MinHash / minimizer distance", "cannot separate domain-sharers"),
    ("SEDEF / BISER", "DNA self-align", "Jaccard + chaining → segdup pairs", "strong on recent SDs; DNA only, no expression"),
    ("Soto 2025 (HSD)", "DNA + copy-number", "SD98 + shared-exon + famCN grouping", "human gold standard; DNA, no RNA/isoform"),
    ("Eichler 2024 (TBC1D3)", "long-read RNA", "map FLNC; assign read iff AS ≥ 10", "ONE family; not de-novo / genome-wide"),
    ("longcallR", "long-read RNA", "CNN SNP + MEC phasing + ASE/ASJ", "uniquely-mappable only; never assigns to COPIES"),
    ("RSEM / Salmon / kallisto", "RNA quant (EM)", "EM apportions multireads", "needs a PRE-DEFINED transcript set; no discovery"),
    ("IsoCon", "long-read RNA", "NN cluster/correct + real-vs-error test", "closest prior art; WITHIN a known family, targeted"),
    ("★ OURS (read-conflict)", "long-read RNA, de-novo", "loci a read CAN'T tell apart (conflict graph) + copy assign", "fills the empty niche: de-novo · genome-wide · RNA · copy-level"),
]


def build():
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def tb(slide, text, l, t, w, h, sz, bold=False, color=(20, 20, 20), align=PP_ALIGN.LEFT, italic=False):
        box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)); tf = box.text_frame; tf.word_wrap = True
        for i, ln in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ln; p.alignment = align
            r = p.runs[0]; r.font.size = Pt(sz); r.font.bold = bold; r.font.italic = italic; r.font.color.rgb = RGBColor(*color)

    def title(slide, s): tb(slide, s, 0.6, 0.28, 12.1, 1.0, 28, bold=True, color=NAVY)

    def fig(slide, png, top=1.5, maxw=11.8, maxh=5.0):
        from PIL import Image
        iw, ih = Image.open(png).size
        sc = min(maxw / (iw / 150), maxh / (ih / 150)); w = (iw / 150) * sc
        slide.shapes.add_picture(png, Inches((13.333 - w) / 2), Inches(top), width=Inches(w))

    # Slide 1
    s = prs.slides.add_slide(blank); title(s, "1 · What is a multi-copy gene family (our definition)?")
    tb(s, "A set of loci the READS cannot tell apart — a family = a connected component of the read-conflict graph "
          "(an edge = some read fits both loci comparably). Relational, annotation-free; no absolute-similarity bar.",
       0.7, 1.15, 11.9, 0.9, 18, italic=True, color=(60, 60, 60))
    fig(s, fig_conflict(), top=2.05, maxh=4.6)

    # Slide 2 — methods table
    s = prs.slides.add_slide(blank); title(s, "2 · Methods to detect multi-copy gene families (gathered)")
    rows, cols = len(METHODS), 4
    gt = s.shapes.add_table(rows, cols, Inches(0.35), Inches(1.35), Inches(12.6), Inches(5.7)).table
    gt.columns[0].width = Inches(2.5); gt.columns[1].width = Inches(2.0); gt.columns[2].width = Inches(4.2); gt.columns[3].width = Inches(3.9)
    for ci in range(cols):
        for ri in range(rows):
            cell = gt.cell(ri, ci); cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.06); cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
            p = cell.text_frame.paragraphs[0]; p.text = METHODS[ri][ci]; p.runs[0].font.size = Pt(11.5 if ri else 12.5)
            header = ri == 0; ours = ri == rows - 1
            p.runs[0].font.bold = header or ours or ci == 0
            if header:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(*NAVY); p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            elif ours:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(225, 240, 225); p.runs[0].font.color.rgb = RGBColor(20, 110, 20)
            else:
                cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if ri % 2 else RGBColor(244, 247, 251)
    tb(s, "The RNA / long-read / de-novo / copy-level niche was empty — IsoCon is the closest, but only WITHIN a known family.",
       0.4, 7.05, 12.5, 0.4, 11.5, italic=True, color=(110, 110, 110))

    # Slide 3 — runaway read
    s = prs.slides.add_slide(blank); title(s, "3 · A runaway read can't wrongly extend the intron chain")
    fig(s, fig_runaway(), top=1.35, maxh=5.5)

    path = os.path.join(OUT, "three_slides.pptx"); prs.save(path); return path


if __name__ == "__main__":
    print("wrote", build())
